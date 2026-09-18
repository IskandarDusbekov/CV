import io
from datetime import timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document

from apps.core.models import SiteSettings
from apps.users.models import PricingPlan, UserProfile, UserSubscription

from .docx_export import THEMES, render_cv_to_docx
from .models import CV, AIUsage
from .services import DEMO_CV_JSON, SUPPORTED_TEMPLATES

User = get_user_model()

META = {"model": "gpt-test", "prompt_tokens": 1000, "completion_tokens": 500, "cost_usd": 0.0012, "duration_ms": 900}
TAILOR_RESULT = (
    {**DEMO_CV_JSON, "job_title": "Backend Engineer"},
    {"match_before": 50, "match_after": 85, "matched_keywords": ["Django"], "missing_keywords": ["Kafka"],
     "changes": ["Summary rewritten"], "vacancy_title": "Backend Engineer", "company": "Acme"},
    META,
)
JOB = "We are hiring a Backend Engineer with Python, Django, PostgreSQL and Kafka experience. " * 2


def make_pro(user):
    plan = PricingPlan.objects.get(scope=PricingPlan.SCOPE_ACCOUNT)
    sub = UserSubscription.objects.create(user=user, profile=user.profile, plan=plan, starts_at=timezone.now() - timedelta(days=1))
    sub.activate()
    return plan


class DownloadPermissionTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="u1")
        self.client.force_login(self.user)

    def _cv(self, template="classic", unlocked=False, **kw):
        return CV.objects.create(user=self.user, raw_input_text="x", cv_json=DEMO_CV_JSON,
                                 selected_template=template, is_unlocked=unlocked, **kw)

    def test_locked_pro_template_cannot_download_any_format(self):
        cv = self._cv("creative")
        for name in ("download_pdf", "download_docx"):
            response = self.client.get(reverse(name, args=[cv.public_id]))
            self.assertEqual(response.status_code, 302)
            self.assertTrue(response["Location"].startswith(reverse("cv_preview", args=[cv.public_id])))

    @mock.patch("apps.cv.views.render_cv_to_pdf", return_value=b"%PDF-1.7")
    def test_first_pdf_is_free_on_free_template_but_word_is_not(self, _):
        first, second = self._cv("ats"), self._cv("simple")
        self.assertEqual(self.client.get(reverse("download_pdf", args=[first.public_id]))["Content-Type"], "application/pdf")
        first.refresh_from_db()
        self.user.profile.refresh_from_db()
        self.assertTrue(first.free_pdf)
        self.assertEqual(self.user.profile.free_pdf_used, 1)
        # shu rezyumeni qayta yuklash — yana bepul, hisob o'smaydi
        self.assertEqual(self.client.get(reverse("download_pdf", args=[first.public_id])).status_code, 200)
        self.user.profile.refresh_from_db()
        self.assertEqual(self.user.profile.free_pdf_used, 1)
        # ikkinchi rezyume — bepul limit tugagan
        self.assertEqual(self.client.get(reverse("download_pdf", args=[second.public_id])).status_code, 302)
        # Word — faqat kredit/Pro
        self.assertEqual(self.client.get(reverse("download_docx", args=[first.public_id])).status_code, 302)
        # Pro shablonga almashtirilsa, bepul PDF ishlamaydi
        first.selected_template = "bold"
        first.save()
        self.assertEqual(self.client.get(reverse("download_pdf", args=[first.public_id])).status_code, 302)

    def test_free_pdf_limit_from_settings_and_template_tiers_from_panel(self):
        from .models import TemplateSetting
        from .services import pdf_access, template_is_pro

        site = SiteSettings.load()
        site.free_pdf_downloads = 0
        site.save()
        self.assertEqual(pdf_access(self.user, self._cv("ats")), "no_free")
        self.assertTrue(template_is_pro("bold"))
        TemplateSetting.objects.filter(code="bold").update(is_pro=False)
        self.assertFalse(template_is_pro("bold"))

    @mock.patch("apps.cv.views.render_cv_to_pdf", return_value=b"%PDF-1.7")
    def test_unlocked_cv_downloads(self, _):
        cv = self._cv("dark", unlocked=True)
        self.assertEqual(self.client.get(reverse("download_pdf", args=[cv.public_id]))["Content-Type"], "application/pdf")
        self.assertIn("wordprocessingml", self.client.get(reverse("download_docx", args=[cv.public_id]))["Content-Type"])

    @mock.patch("apps.cv.views.render_cv_to_pdf", return_value=b"%PDF-1.7")
    def test_tailored_copy_of_unlocked_cv_is_unlocked(self, _):
        root = self._cv(unlocked=True)
        copy = self._cv(parent=root)
        self.assertEqual(self.client.get(reverse("download_pdf", args=[copy.public_id])).status_code, 200)

    @mock.patch("apps.cv.views.render_cv_to_pdf", return_value=b"%PDF-1.7")
    def test_pro_user_downloads_everything(self, _):
        make_pro(self.user)
        self.assertEqual(self.client.get(reverse("download_pdf", args=[self._cv().public_id])).status_code, 200)

    def test_anonymous_download_redirects_to_login(self):
        self.client.logout()
        cv = CV.objects.create(raw_input_text="x", cv_json=DEMO_CV_JSON, is_unlocked=True)
        session = self.client.session
        session["owned_cv_ids"] = [cv.id]
        session.save()
        self.assertTrue(self.client.get(reverse("download_pdf", args=[cv.public_id]))["Location"].startswith(reverse("user_login")))

    def test_other_users_cv_is_404(self):
        cv = CV.objects.create(user=User.objects.create_user(username="u2"), raw_input_text="x", cv_json=DEMO_CV_JSON)
        self.assertEqual(self.client.get(reverse("cv_preview", args=[cv.public_id])).status_code, 404)


@mock.patch("apps.cv.views.generate_cv_from_text", return_value=(dict(DEMO_CV_JSON), META))
class GenerateQuotaTests(TestCase):
    def _generate(self, client=None):
        return (client or self.client).post(reverse("generate_cv"), {"text": "men sardor, 6 yil python dasturchiman"})

    def test_anonymous_gets_two_free_cvs(self, _):
        self.assertEqual(self._generate().status_code, 200)
        self.assertEqual(self._generate().status_code, 200)
        response = self._generate()
        self.assertEqual(response.status_code, 403)
        self.assertTrue(response.json()["limit_reached"])

    def test_clearing_cookies_does_not_reset_limit_same_ip(self, _):
        self._generate(); self._generate()
        fresh = self.client_class()
        self.assertEqual(self._generate(fresh).status_code, 403)

    def test_pro_uses_plan_limit(self, _):
        user = User.objects.create_user(username="pro")
        plan = make_pro(user)
        plan.max_cvs = 3
        plan.save()
        self.client.force_login(user)
        codes = [self._generate().status_code for _ in range(4)]
        self.assertEqual(codes, [200, 200, 200, 403])


@mock.patch("apps.cv.views.tailor_cv_to_job", return_value=TAILOR_RESULT)
class TailorTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="u1")
        self.client.force_login(self.user)
        self.cv = CV.objects.create(user=self.user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template="ats")

    def _tailor(self, cv=None):
        return self.client.post(reverse("tailor_cv", args=[(cv or self.cv).public_id]), {"job_description": JOB})

    def test_creates_new_version_and_keeps_original(self, _):
        response = self._tailor()
        copy = CV.objects.exclude(id=self.cv.id).get()
        self.assertRedirects(response, reverse("cv_preview", args=[copy.public_id]), fetch_redirect_response=False)
        self.assertEqual(copy.parent, self.cv)
        self.assertEqual(copy.tailor_report["match_after"], 85)
        self.cv.refresh_from_db()
        self.assertEqual(self.cv.cv_json["job_title"], DEMO_CV_JSON["job_title"])
        self.assertContains(self.client.get(reverse("cv_preview", args=[copy.public_id])), "Kafka")

    def test_free_limit_one(self, _):
        self._tailor(); self._tailor()
        self.assertEqual(CV.objects.count(), 2)

    def test_unlocked_cv_gets_site_tailorings(self, _):
        site = SiteSettings.load()
        site.tailor_per_unlocked_cv = 2
        site.save()
        self.cv.unlock()
        for _ in range(4):
            self._tailor()
        self.assertEqual(AIUsage.objects.filter(kind=AIUsage.KIND_TAILOR).count(), 2)

    def test_short_job_description_rejected(self, mocked):
        self.client.post(reverse("tailor_cv", args=[self.cv.public_id]), {"job_description": "python"})
        mocked.assert_not_called()


class DocxExportTests(TestCase):
    def test_every_template_has_theme_and_renders(self):
        self.assertEqual(set(THEMES), SUPPORTED_TEMPLATES)
        for code in SUPPORTED_TEMPLATES:
            cv = CV(raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template=code, is_unlocked=True)
            doc = Document(io.BytesIO(render_cv_to_docx(cv, None)))
            text = "\n".join(p.text for p in doc.paragraphs)
            text += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn("Sardor Nazarov", text, code)
            self.assertIn("Uzum Market", text, code)

    def test_ats_docx_has_no_tables(self):
        for code in ("ats", "ats_modern"):
            cv = CV(raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template=code, is_unlocked=True)
            self.assertEqual(len(Document(io.BytesIO(render_cv_to_docx(cv, None))).tables), 0, code)


class PageRenderTests(TestCase):
    def test_public_pages(self):
        for url in ["/", "/pricing/", "/cv/templates/", "/cv/namuna/", "/cv/builder/", "/users/login/", "/users/login/?mode=signup",
                    "/namunalar/", "/namunalar/kassir/", "/sitemap.xml", "/qollanma/diplomsiz-ish-topsa-boladimi/",
                    *[f"/cv/template-preview/{c}/" for c in SUPPORTED_TEMPLATES]]:
            self.assertEqual(self.client.get(url).status_code, 200, url)

    def test_landing_has_auth_links_and_sample_cta(self):
        response = self.client.get("/")
        self.assertContains(response, 'href="/users/login/"')
        self.assertContains(response, 'href="/users/signup/"')
        self.assertContains(response, 'href="/cv/namuna/"')

    def test_owner_pages(self):
        user = User.objects.create_user(username="u1")
        self.client.force_login(user)
        cv = CV.objects.create(user=user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template="creative")
        response = self.client.get(reverse("cv_preview", args=[cv.public_id]))
        self.assertContains(response, "Yuklab olish uchun oching")
        self.assertContains(response, "Vakansiyaga moslashtirish")
        self.assertEqual(self.client.get(reverse("user_dashboard")).status_code, 200)


class CreditUnlockTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="u1")
        self.client.force_login(self.user)
        self.cv = CV.objects.create(user=self.user, raw_input_text="x", cv_json=DEMO_CV_JSON)

    def test_unlock_spends_one_credit(self):
        UserProfile.objects.filter(user=self.user).update(credits=2)
        self.client.post(reverse("unlock_cv", args=[self.cv.public_id]))
        self.cv.refresh_from_db()
        self.assertTrue(self.cv.is_unlocked)
        self.assertEqual(UserProfile.objects.get(user=self.user).credits, 1)
        # qayta bosish kredit yechmaydi
        self.client.post(reverse("unlock_cv", args=[self.cv.public_id]))
        self.assertEqual(UserProfile.objects.get(user=self.user).credits, 1)

    def test_no_credits_no_unlock(self):
        self.client.post(reverse("unlock_cv", args=[self.cv.public_id]))
        self.cv.refresh_from_db()
        self.assertFalse(self.cv.is_unlocked)

    def test_numeric_id_urls_do_not_exist(self):
        self.assertEqual(self.client.get(f"/cv/preview/{self.cv.id}/").status_code, 404)
        self.assertEqual(self.client.get(f"/cv/preview/{self.cv.public_id}/").status_code, 200)

    def test_staff_can_view_any_cv(self):
        admin = User.objects.create_superuser("adm", "a@a.uz", "x12345678")
        self.client.force_login(admin)
        self.assertContains(self.client.get(reverse("cv_preview", args=[self.cv.public_id])), "Admin ko")


@mock.patch("apps.cv.views.render_cv_to_pdf", return_value=b"%PDF-1.7")
class TelegramAppDownloadTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="tg")
        UserProfile.objects.filter(user=self.user).update(telegram_id=4242)
        self.client.force_login(self.user)
        self.cv = CV.objects.create(user=self.user, raw_input_text="x", cv_json=DEMO_CV_JSON, is_unlocked=True)

    def _link(self, fmt="pdf", cv=None):
        return self.client.post(reverse("download_link", args=[(cv or self.cv).public_id, fmt]))

    def test_signed_link_downloads_without_session(self, _):
        data = self._link().json()
        self.assertTrue(data["file_name"].endswith(".pdf"))
        anon = self.client_class()
        r = anon.get(data["url"].replace("http://testserver", ""))
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertIn("attachment", r["Content-Disposition"])
        self.assertEqual(r["Access-Control-Allow-Origin"], "https://web.telegram.org")
        docx = anon.get(self._link("docx").json()["url"].replace("http://testserver", ""))
        self.assertIn("wordprocessingml", docx["Content-Type"])

    def test_tampered_or_expired_link_rejected(self, _):
        url = self._link().json()["url"].replace("http://testserver", "")
        self.assertEqual(self.client_class().get(url[:-3] + "abc/").status_code, 410)
        with mock.patch("apps.cv.views._DL_MAX_AGE", -1):
            self.assertEqual(self.client_class().get(url).status_code, 410)

    def test_locked_or_foreign_cv_gets_no_link(self, _):
        locked = CV.objects.create(user=self.user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template="dark")
        self.assertEqual(self._link(cv=locked).status_code, 403)
        self.assertEqual(self._link("docx", cv=CV.objects.create(user=self.user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template="ats")).status_code, 403)
        foreign = CV.objects.create(user=User.objects.create_user(username="z"), raw_input_text="x", cv_json=DEMO_CV_JSON, is_unlocked=True)
        self.assertEqual(self._link(cv=foreign).status_code, 404)
        self.client.logout()
        self.assertEqual(self._link().status_code, 401)

    @mock.patch("apps.users.bot.send_document", return_value=True)
    def test_send_to_telegram_chat(self, send, _):
        r = self.client.post(reverse("send_to_telegram", args=[self.cv.public_id, "pdf"]))
        self.assertEqual(r.json(), {"ok": True})
        chat_id, content, filename = send.call_args.args[:3]
        self.assertEqual((chat_id, content), (4242, b"%PDF-1.7"))
        self.assertTrue(filename.endswith(".pdf"))

    def test_telegram_script_only_inside_mini_app(self, _):
        preview = reverse("cv_preview", args=[self.cv.public_id])
        self.assertNotContains(self.client.get(preview), "telegram-web-app.js")
        session = self.client.session
        session["in_telegram"] = True
        session.save()
        r = self.client.get(preview)
        self.assertContains(r, "telegram-web-app.js")
        self.assertContains(r, f'data-dl="{self.cv.public_id}:docx"')


class MissingDetailsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="d", first_name="Dilnoza", last_name="Karimova")
        UserProfile.objects.filter(user=self.user).update(phone="+998901112233")
        self.client.force_login(self.user)
        data = {**DEMO_CV_JSON, "phone": "", "email": "", "location": ""}
        self.cv = CV.objects.create(user=self.user, raw_input_text="x", cv_json=data)
        self.child = CV.objects.create(user=self.user, raw_input_text="x", cv_json={**data, "job_title": "Backend"}, parent=self.cv)

    def test_preview_asks_only_missing_fields_prefilled_from_profile(self):
        missing = self.client.get(reverse("cv_preview", args=[self.cv.public_id])).context["missing_details"]
        self.assertEqual([f["name"] for f in missing], ["phone", "email", "location"])
        self.assertEqual(missing[0]["value"], "+998901112233")

    def test_save_updates_whole_family_without_touching_tailored_title(self):
        self.client.post(reverse("cv_details", args=[self.child.public_id]),
                         {"action": "save", "phone": "+998901112233", "location": " Toshkent ", "job_title": "Dev"})
        self.cv.refresh_from_db(); self.child.refresh_from_db()
        self.assertEqual((self.cv.cv_json["phone"], self.cv.cv_json["location"]), ("+998901112233", "Toshkent"))
        self.assertEqual(self.child.cv_json["location"], "Toshkent")
        self.assertEqual(self.cv.cv_json["job_title"], DEMO_CV_JSON["job_title"])
        self.assertEqual(self.child.cv_json["job_title"], "Dev")

    def test_skip_hides_card(self):
        self.client.post(reverse("cv_details", args=[self.cv.public_id]), {"action": "skip"})
        self.assertEqual(self.client.get(reverse("cv_preview", args=[self.cv.public_id])).context["missing_details"], [])

    def test_other_user_cannot_edit(self):
        self.client.force_login(User.objects.create_user(username="x"))
        self.assertEqual(self.client.post(reverse("cv_details", args=[self.cv.public_id]), {"phone": "1"}).status_code, 404)

    def test_builder_prefills_telegram_name_and_phone(self):
        r = self.client.get(reverse("cv_builder"))
        self.assertContains(r, 'value="Dilnoza Karimova"')
        self.assertContains(r, 'value="+998901112233"')


@mock.patch("apps.cv.views.generate_cv_from_text", return_value=(dict(DEMO_CV_JSON), META))
class AICostTests(TestCase):
    def test_usage_cost_and_activity_recorded(self, _):
        self.client.post(reverse("generate_cv"), {"text": "men sardor, 6 yil python dasturchiman"})
        usage = AIUsage.objects.get()
        self.assertEqual((usage.prompt_tokens, usage.completion_tokens, usage.model), (1000, 500, "gpt-test"))
        self.assertGreater(usage.cost_usd, 0)
        from apps.core.models import ActivityLog
        self.assertTrue(ActivityLog.objects.filter(action="cv_create").exists())

    def test_failed_ai_call_is_logged_but_not_counted(self, mocked):
        from .services import AIError
        mocked.side_effect = AIError("timeout", {"model": "gpt-test", "duration_ms": 30000})
        for _ in range(3):
            self.assertEqual(self.client.post(reverse("generate_cv"), {"text": "men sardor, 6 yil python dasturchiman"}).status_code, 502)
        self.assertEqual(AIUsage.objects.filter(success=False).count(), 3)
        mocked.side_effect = None
        self.assertEqual(self.client.post(reverse("generate_cv"), {"text": "men sardor, 6 yil python dasturchiman"}).status_code, 200)


class SamplesAndEditorTests(TestCase):
    def test_sample_to_editor_flow_for_anonymous(self):
        from .models import ResumeSample

        response = self.client.post("/namunalar/kassir/boshlash/")
        cv = CV.objects.get()
        self.assertRedirects(response, f"/cv/edit/{cv.public_id}/?new=1", fetch_redirect_response=False)
        self.assertEqual(cv.cv_json["full_name"], "Malika Tursunova")
        self.assertFalse(AIUsage.objects.exists())  # AI ishlatilmaydi
        self.assertEqual(ResumeSample.objects.get(slug="kassir").uses, 1)
        self.assertContains(self.client.get(f"/cv/edit/{cv.public_id}/?new=1"), "Namuna nusxalandi")

        response = self.client.post(f"/cv/edit/{cv.public_id}/", {
            "full_name": "Ali Valiyev", "job_title": "Kassir", "phone": "+998 91 000 00 00", "skills": "Excel, 1C\nKassa",
            "languages": "O'zbek — ona tili", "exp-0-position": "Kassir", "exp-0-company": "Havas", "exp-0-duration": "2024 — hozir",
            "exp-0-responsibilities": "Kassada ishladim\n• Hisobot topshirdim", "exp-7-position": "", "edu-3-institution": "Kollej",
            "edu-3-degree": "Buxgalteriya", "edu-3-year": "2023",
        })
        self.assertRedirects(response, reverse("cv_preview", args=[cv.public_id]), fetch_redirect_response=False)
        cv.refresh_from_db()
        self.assertEqual(cv.cv_json["full_name"], "Ali Valiyev")
        self.assertEqual(cv.cv_json["skills"], ["Excel", "1C", "Kassa"])
        self.assertEqual(cv.cv_json["experience"], [{"position": "Kassir", "company": "Havas", "duration": "2024 — hozir",
                                                     "responsibilities": ["Kassada ishladim", "Hisobot topshirdim"]}])
        self.assertEqual(cv.cv_json["education"][0]["institution"], "Kollej")
        self.assertEqual(cv.cv_json["_language"], "uz")

    def test_logged_in_user_sample_uses_own_name_and_other_user_cannot_edit(self):
        user = User.objects.create_user(username="s", first_name="Sardor", last_name="Aliyev")
        UserProfile.objects.filter(user=user).update(phone="+998901112233")
        self.client.force_login(user)
        self.client.post("/namunalar/haydovchi/boshlash/")
        cv = CV.objects.get()
        self.assertEqual((cv.cv_json["full_name"], cv.cv_json["phone"], cv.user_id), ("Sardor Aliyev", "+998901112233", user.pk))
        self.client.force_login(User.objects.create_user(username="other"))
        self.assertEqual(self.client.get(f"/cv/edit/{cv.public_id}/").status_code, 404)

    def test_renaming_free_pdf_cv_resets_free_pdf(self):
        user = User.objects.create_user(username="r")
        self.client.force_login(user)
        cv = CV.objects.create(user=user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template="ats", free_pdf=True)
        self.client.post(f"/cv/edit/{cv.public_id}/", {"full_name": "Boshqa Odam"})
        cv.refresh_from_db()
        self.assertFalse(cv.free_pdf)

    def test_preview_shows_template_strip_with_tiers_and_referral(self):
        user = User.objects.create_user(username="p")
        self.client.force_login(user)
        cv = CV.objects.create(user=user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template="ats")
        response = self.client.get(reverse("cv_preview", args=[cv.public_id]))
        self.assertContains(response, 'id="tpl-strip"')
        self.assertContains(response, "PDF yuklab olish — bepul")
        self.assertContains(response, '<em class="pro">PRO</em>')
        self.assertContains(response, "/r/")


def editor_post(data, **extra):
    """Rezyume JSON'ini tahrirlash formasi yuboradigan POST ko'rinishiga aylantiradi."""
    post = {k: data.get(k, "") for k in ("full_name", "job_title", "email", "phone", "location", "github", "linkedin", "summary")}
    post["skills"] = "\n".join(data.get("skills", []))
    post["languages"] = "\n".join(data.get("languages", []))
    for i, e in enumerate(data.get("experience", [])):
        post.update({f"exp-{i}-position": e.get("position", ""), f"exp-{i}-company": e.get("company", ""),
                     f"exp-{i}-duration": e.get("duration", ""), f"exp-{i}-responsibilities": "\n".join(e.get("responsibilities", []))})
    for i, e in enumerate(data.get("education", [])):
        post.update({f"edu-{i}-institution": e.get("institution", ""), f"edu-{i}-degree": e.get("degree", ""), f"edu-{i}-year": e.get("year", "")})
    post.update(extra)
    return post


IMPROVED = {**DEMO_CV_JSON, "full_name": "Aziza Tursunova", "job_title": "Kassir"}


class ImproveWithAITests(TestCase):
    def _sample_cv(self):
        self.client.post("/namunalar/kassir/boshlash/")
        return CV.objects.get()

    @mock.patch("apps.cv.services.improve_cv", return_value=(IMPROVED, META))
    def test_sample_needs_notes_then_free_user_gets_one_ai_fill(self, mocked):
        cv = self._sample_cv()
        page = self.client.get(f"/cv/edit/{cv.public_id}/?new=1")
        self.assertContains(page, "AI bilan to'ldirish")
        self.assertContains(page, "1 ta bepul")

        # namuna o'zgarmagan va izoh yo'q — AI chaqirilmaydi, limit ham ketmaydi
        url = reverse("cv_improve", args=[cv.public_id])
        self.client.post(url, editor_post(cv.cv_json))
        mocked.assert_not_called()
        self.assertFalse(AIUsage.objects.exists())

        notes = "Korzinkada 2 yil kassir bo'ldim, hozir Makroda sotuvchiman. Excel bilaman."
        response = self.client.post(url, editor_post(cv.cv_json, ai_notes=notes))
        self.assertRedirects(response, f"/cv/edit/{cv.public_id}/?ai=1", fetch_redirect_response=False)
        args, kwargs = mocked.call_args
        self.assertEqual(args[1], notes)
        self.assertEqual(kwargs["sample_json"]["full_name"], "Malika Tursunova")  # AI namunadagi faktlarni ajratadi
        cv.refresh_from_db()
        self.assertEqual(cv.cv_json["full_name"], "Aziza Tursunova")
        self.assertEqual(AIUsage.objects.get().kind, AIUsage.KIND_IMPROVE)
        self.assertContains(self.client.get(f"/cv/edit/{cv.public_id}/?ai=1"), "AI rezyumeni to'ldirdi")

        # bepul limit — 1 marta; qo'lda tahrirlash ishlayveradi
        self.client.post(url, editor_post(cv.cv_json, ai_notes=notes))
        self.assertEqual(mocked.call_count, 1)
        page = self.client.get(f"/cv/edit/{cv.public_id}/")
        self.assertContains(page, "Bepul AI to&#x27;ldirish ishlatildi")
        self.assertNotContains(page, f'formaction="{url}"')

    @mock.patch("apps.cv.services.improve_cv")
    def test_ai_error_is_not_counted_and_keeps_edits_and_notes(self, mocked):
        from .services import AIError

        mocked.side_effect = AIError("timeout", META)
        user = User.objects.create_user(username="e")
        self.client.force_login(user)
        cv = CV.objects.create(user=user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template="ats")
        url = reverse("cv_improve", args=[cv.public_id])
        self.client.post(url, editor_post(DEMO_CV_JSON, full_name="Yangi Ism", ai_notes="IELTS 7.0 oldim"))
        cv.refresh_from_db()
        self.assertEqual(cv.cv_json["full_name"], "Yangi Ism")  # formadagi o'zgarish yo'qolmadi
        self.assertEqual(AIUsage.objects.get().success, False)
        page = self.client.get(f"/cv/edit/{cv.public_id}/")
        self.assertContains(page, "IELTS 7.0 oldim")
        self.assertContains(page, "1 ta bepul")
        self.assertIsNone(mocked.call_args.kwargs["sample_json"])  # oddiy rezyume — namuna emas

    @mock.patch("apps.cv.services.improve_cv", return_value=(IMPROVED, META))
    def test_unlocked_cv_gets_its_own_limit_and_other_user_is_404(self, mocked):
        user = User.objects.create_user(username="u")
        self.client.force_login(user)
        AIUsage.objects.create(user=user, kind=AIUsage.KIND_IMPROVE)  # bepul limit ishlatilgan
        cv = CV.objects.create(user=user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template="ats", is_unlocked=True)
        url = reverse("cv_improve", args=[cv.public_id])
        for _ in range(4):
            self.client.post(url, editor_post(DEMO_CV_JSON))
        self.assertEqual(mocked.call_count, SiteSettings.load().improve_per_unlocked_cv)

        self.client.force_login(User.objects.create_user(username="x"))
        self.assertEqual(self.client.post(url, editor_post(DEMO_CV_JSON)).status_code, 404)


class LuckyGiftTests(TestCase):
    def setUp(self):
        from .models import LuckyGift

        self.gift = LuckyGift.load()
        self.gift.is_active = True
        self.gift.save()
        self.user = User.objects.create_user(username="lucky", first_name="Aziza")
        self.client.force_login(self.user)

    def _cv(self, template="bold", **kw):
        return CV.objects.create(user=self.user, raw_input_text="x", cv_json=DEMO_CV_JSON, selected_template=template, **kw)

    @mock.patch("apps.cv.views.render_cv_to_pdf", return_value=b"%PDF-1.7")
    def test_gift_opens_any_template_with_watermark_but_not_word(self, _):
        from .models import LuckyGrant

        cv = self._cv("bold")  # Pro shablon — sovg'asiz yopiq bo'lardi
        page = self.client.get(reverse("cv_preview", args=[cv.public_id]))
        self.assertContains(page, self.gift.title)
        self.assertContains(page, "Men ham xursandman")
        grant = LuckyGrant.objects.get()
        self.assertEqual((grant.user_id, grant.cv_id, grant.downloaded), (self.user.pk, cv.pk, False))

        self.assertEqual(self.client.get(reverse("download_pdf", args=[cv.public_id]))["Content-Type"], "application/pdf")
        cv.refresh_from_db()
        self.user.profile.refresh_from_db()
        grant.refresh_from_db()
        self.assertTrue(grant.downloaded)
        self.assertEqual(self.user.profile.free_pdf_used, 0)  # bepul PDF saqlanib qoladi
        self.assertFalse(cv.free_pdf)
        self.assertEqual(self.client.get(reverse("download_docx", args=[cv.public_id])).status_code, 302)  # Word baribir yopiq

        from .services import build_cv_context

        self.assertTrue(build_cv_context(cv, self.user)["show_watermark"])  # sovg'ada sayt belgisi doim bor

        # Ikkinchi rezyumega sovg'a berilmaydi
        other = self._cv("dark")
        self.client.get(reverse("cv_preview", args=[other.public_id]))
        self.assertEqual(LuckyGrant.objects.count(), 1)
        self.assertEqual(self.client.get(reverse("download_pdf", args=[other.public_id])).status_code, 302)

    def test_reaction_is_saved_once(self):
        cv = self._cv()
        self.client.get(reverse("cv_preview", args=[cv.public_id]))
        url = reverse("lucky_reaction", args=[cv.public_id])
        self.assertEqual(self.client.post(url, {"value": "🙏 Rahmat"}).json()["thanks"], self.gift.thanks_text)
        self.client.post(url, {"value": "😊 Men ham xursandman"})
        self.assertEqual(self.client.post(url, {"value": "boshqa"}).json()["ok"], False)
        from .models import LuckyGrant

        self.assertEqual(LuckyGrant.objects.get().reaction, "🙏 Rahmat")

    def test_switched_off_audience_and_daily_limit(self):
        from .models import LuckyGift, LuckyGrant

        self.gift.is_active = False
        self.gift.save()
        cv = self._cv()
        self.client.get(reverse("cv_preview", args=[cv.public_id]))
        self.assertFalse(LuckyGrant.objects.exists())

        self.gift.is_active = True
        self.gift.audience = LuckyGift.AUDIENCE_NOT_PAID
        self.gift.save()
        UserProfile.objects.filter(user=self.user).update(credits=2)
        self.client.get(reverse("cv_preview", args=[cv.public_id]))
        self.assertFalse(LuckyGrant.objects.exists())  # to'lagan odamga sovg'a yo'q

        UserProfile.objects.filter(user=self.user).update(credits=0)
        self.gift.daily_limit = 1
        self.gift.save()
        other = User.objects.create_user(username="second")
        LuckyGrant.objects.create(user=other, cv=CV.objects.create(user=other, raw_input_text="x", cv_json=DEMO_CV_JSON))
        self.client.get(reverse("cv_preview", args=[cv.public_id]))
        self.assertEqual(LuckyGrant.objects.filter(user=self.user).count(), 0)  # kunlik chegara to'lgan

        self.gift.daily_limit = 5
        self.gift.save()
        self.assertContains(self.client.get(reverse("cv_preview", args=[cv.public_id])), self.gift.title)


class InterfaceTests(TestCase):
    def test_icon_tag_and_cabinet_button(self):
        from apps.core.templatetags.ui import icon

        self.assertIn("<svg", icon("edit"))
        self.assertEqual(icon("no-such-icon"), "")
        self.assertNotContains(self.client.get("/cv/builder/"), 'aria-label="Kabinet"')
        self.client.force_login(User.objects.create_user(username="k"))
        self.assertContains(self.client.get("/cv/builder/"), 'aria-label="Kabinet"')


class SeoTests(TestCase):
    def test_sitemap_robots_and_meta(self):
        from apps.core.models import SeoPage

        sitemap = self.client.get("/sitemap.xml").content.decode()
        for path in ("/namunalar/kassir/", "/qollanma/", "/cv/template-preview/bold/"):
            self.assertIn(path, sitemap)
        self.assertIn("Sitemap:", self.client.get("/robots.txt").content.decode())
        home = self.client.get("/").content.decode()
        self.assertIn('rel="canonical"', home)
        self.assertIn('"@type": "SoftwareApplication"', home)

        SeoPage.objects.create(path="namunalar", title="Maxsus sarlavha", description="Maxsus tavsif", noindex=True)
        page = self.client.get("/namunalar/").content.decode()
        self.assertIn("<title>Maxsus sarlavha</title>", page)
        self.assertIn('content="Maxsus tavsif"', page)
        self.assertIn('content="noindex, follow"', page)

    def test_google_html_file_and_meta_verification_from_panel(self):
        from apps.core.models import SiteSettings

        self.assertEqual(self.client.get("/google1a2b3c4d5e6f7a8b.html").status_code, 404)  # sozlanmagan
        admin = User.objects.create_user(username="seo", is_staff=True, is_superuser=True)
        self.client.force_login(admin)
        form = self.client.get(reverse("panel:seo")).context["settings_form"]
        data = {"section": "settings", **{k: (v if v is not None else "") for k, v in form.initial.items() if k in form.fields}}
        data["google_verification_file"] = "google-site-verification: google1a2b3c4d5e6f7a8b.html"
        data["google_site_verification"] = '<meta name="google-site-verification" content="AbC-123_x" />'
        self.client.post(reverse("panel:seo"), data)
        site = SiteSettings.objects.get()
        self.assertEqual((site.google_verification_file, site.google_site_verification), ("google1a2b3c4d5e6f7a8b.html", "AbC-123_x"))

        self.client.logout()
        response = self.client.get("/google1a2b3c4d5e6f7a8b.html")
        self.assertEqual(response.content.decode(), "google-site-verification: google1a2b3c4d5e6f7a8b.html")
        self.assertEqual(self.client.get("/google0000000000000000.html").status_code, 404)  # boshqa nom
        self.assertContains(self.client.get("/"), '<meta name="google-site-verification" content="AbC-123_x">')

        data["google_verification_file"] = "noto'g'ri qiymat"
        self.client.force_login(admin)
        self.assertContains(self.client.post(reverse("panel:seo"), data), "Google bergan fayl nomini yozing")

    def test_private_pages_noindex_and_guide_question_page(self):
        self.assertContains(self.client.get("/users/login/"), 'content="noindex, nofollow"')
        response = self.client.get("/qollanma/diplomsiz-ish-topsa-boladimi/")
        self.assertContains(response, "<h1>Diplomsiz ish topsa bo&#x27;ladimi?</h1>", html=False)
        self.assertContains(response, '"@type": "FAQPage"')
        self.assertEqual(self.client.get("/qollanma/yoq-savol/").status_code, 404)
